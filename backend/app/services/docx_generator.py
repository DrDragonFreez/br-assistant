"""Формирование черновиков документов в формате DOCX.

Работает в двух режимах:

* **По бланку.** Если файл шаблона есть, он заполняется через docxtpl —
  оформление, бланк и реквизиты организации сохраняются полностью.
* **Структурная сборка.** Если файла нет, документ собирается python-docx
  по схеме полей: ГОСТ-подобные поля страницы, Times New Roman 14,
  реквизиты в нужном порядке. Оформление проще, но документ пригоден к работе.

Второй режим нужен, потому что на хакатоне исходные бланки могут прийти
позже кода, а демонстрация генерации должна работать в любом случае.

Генерация синхронная (python-docx не умеет иначе), поэтому выполняется
в пуле потоков и вызывается из Celery для больших документов.
"""

from __future__ import annotations

import asyncio
import re
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from app.core.config import settings
from app.core.exceptions import DocumentParseError
from app.core.logging import get_logger
from app.core.metrics import documents_generated_total

logger = get_logger(__name__)

UNSAFE_FILENAME_RE = re.compile(r"[^\w\s\-.()а-яёА-ЯЁ]", re.UNICODE)

MONTHS_RU: tuple[str, ...] = (
    "января",
    "февраля",
    "марта",
    "апреля",
    "мая",
    "июня",
    "июля",
    "августа",
    "сентября",
    "октября",
    "ноября",
    "декабря",
)


def format_date_ru(value: datetime | None = None) -> str:
    """Форматирует дату по правилам делопроизводства.

    Args:
        value: Дата; по умолчанию текущая.

    Returns:
        Строка вида «15 марта 2026 г.».
    """
    date = value or datetime.now(UTC)
    return f"{date.day} {MONTHS_RU[date.month - 1]} {date.year} г."


def safe_filename(name: str, *, max_length: int = 120) -> str:
    """Приводит имя файла к безопасному виду.

    Args:
        name: Исходное имя.
        max_length: Ограничение длины.

    Returns:
        Имя без служебных символов.
    """
    cleaned = UNSAFE_FILENAME_RE.sub("_", name).strip(" ._")
    cleaned = re.sub(r"_{2,}", "_", cleaned)
    return (cleaned or "document")[:max_length]


def _normalize_value(value: Any) -> str:
    """Приводит значение поля к строке для вставки в документ.

    Args:
        value: Значение поля любого типа.

    Returns:
        Строковое представление.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, (list, tuple)):
        return "\n".join(f"{index}. {item}" for index, item in enumerate(value, start=1))
    if isinstance(value, datetime):
        return format_date_ru(value)
    return str(value)


class DocxGenerator:
    """Генератор черновиков документов."""

    def __init__(self, output_dir: Path | None = None) -> None:
        """Инициализирует генератор.

        Args:
            output_dir: Каталог для сформированных файлов.
        """
        self.output_dir = Path(output_dir or settings.ARTIFACTS_DIR) / "documents"
        self.output_dir.mkdir(parents=True, exist_ok=True)

    # ─────────────────── Режим 1: заполнение бланка ───────────────────
    def _render_from_template_sync(
        self,
        template_path: Path,
        context: dict[str, Any],
        output_path: Path,
    ) -> None:
        """Заполняет готовый бланк через docxtpl (синхронно).

        Args:
            template_path: Путь к файлу бланка.
            context: Значения плейсхолдеров.
            output_path: Куда сохранить результат.

        Raises:
            DocumentParseError: Бланк повреждён или содержит ошибку в разметке.
        """
        try:
            from docxtpl import DocxTemplate

            document = DocxTemplate(str(template_path))
            document.render({key: _normalize_value(value) for key, value in context.items()})
            document.save(str(output_path))
        except Exception as exc:  # noqa: BLE001
            raise DocumentParseError(
                "Не удалось заполнить бланк документа",
                context={"template": template_path.name, "error": str(exc)},
            ) from exc

    # ─────────────────── Режим 2: структурная сборка ───────────────────
    def _build_structured_sync(
        self,
        *,
        title: str,
        fields_schema: list[dict[str, Any]],
        values: dict[str, Any],
        approvals: list[dict[str, Any]],
        related_clauses: list[str],
        output_path: Path,
    ) -> None:
        """Собирает документ по схеме полей (синхронно).

        Args:
            title: Название документа.
            fields_schema: Схема полей шаблона.
            values: Собранные значения.
            approvals: Маршрут согласования.
            related_clauses: Пункты инструкции, определяющие требования.
            output_path: Куда сохранить результат.
        """
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Cm, Pt

        document = Document()

        section = document.sections[0]
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)

        normal = document.styles["Normal"]
        normal.font.name = "Times New Roman"
        normal.font.size = Pt(14)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.first_line_indent = Cm(1.25)

        header = document.add_paragraph()
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header.paragraph_format.first_line_indent = Cm(0)
        header_run = header.add_run("ЦЕНТРАЛЬНЫЙ БАНК РОССИЙСКОЙ ФЕДЕРАЦИИ")
        header_run.bold = True

        if values.get("department"):
            subheader = document.add_paragraph()
            subheader.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subheader.paragraph_format.first_line_indent = Cm(0)
            subheader.add_run(_normalize_value(values["department"]))

        document.add_paragraph()

        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.first_line_indent = Cm(0)
        heading_run = heading.add_run(title.upper())
        heading_run.bold = True
        heading_run.font.size = Pt(16)

        requisites = document.add_paragraph()
        requisites.paragraph_format.first_line_indent = Cm(0)
        doc_number = values.get("doc_number") or "№ ____________"
        requisites.add_run(f"{format_date_ru()}\t\t\t\t{doc_number}")

        addressee = values.get("addressee") or values.get("recipient_org")
        if addressee:
            document.add_paragraph()
            block = document.add_paragraph()
            block.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            block.paragraph_format.first_line_indent = Cm(0)
            block.add_run(_normalize_value(addressee))
            if values.get("recipient_person"):
                block.add_run(f"\n{_normalize_value(values['recipient_person'])}")

        if values.get("subject") or values.get("title"):
            document.add_paragraph()
            subject = document.add_paragraph()
            subject.paragraph_format.first_line_indent = Cm(0)
            subject.add_run("О ")
            subject.add_run(_normalize_value(values.get("subject") or values.get("title")))

        document.add_paragraph()

        skip_keys = {
            "department",
            "doc_number",
            "addressee",
            "recipient_org",
            "recipient_person",
            "subject",
            "title",
            "author_name",
            "author_position",
            "signer",
        }
        for field in fields_schema:
            name = str(field.get("name", ""))
            if name in skip_keys:
                continue
            value = values.get(name)
            if not value:
                continue

            label = str(field.get("label", name))
            if field.get("type") == "list" and isinstance(value, (list, tuple)):
                caption = document.add_paragraph()
                caption.paragraph_format.first_line_indent = Cm(0)
                caption.add_run(f"{label}:").bold = True
                for item in value:
                    document.add_paragraph(str(item), style="List Number")
            elif field.get("type") == "text":
                document.add_paragraph(_normalize_value(value))
            else:
                paragraph = document.add_paragraph()
                paragraph.add_run(f"{label}: ").bold = True
                paragraph.add_run(_normalize_value(value))

        document.add_paragraph()
        document.add_paragraph()
        signature = document.add_paragraph()
        signature.paragraph_format.first_line_indent = Cm(0)
        position = values.get("author_position") or values.get("signer") or "Должность"
        name = values.get("author_name") or "И.О. Фамилия"
        signature.add_run(
            f"{_normalize_value(position)}\t\t_____________\t\t{_normalize_value(name)}"
        )

        if approvals:
            document.add_page_break()
            approval_heading = document.add_paragraph()
            approval_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            approval_heading.paragraph_format.first_line_indent = Cm(0)
            approval_heading.add_run("ЛИСТ СОГЛАСОВАНИЯ").bold = True
            document.add_paragraph()

            table = document.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            headers = ("№", "Подразделение", "Срок, раб. дн.", "Виза, дата")
            for cell, text in zip(table.rows[0].cells, headers, strict=True):
                cell.paragraphs[0].paragraph_format.first_line_indent = Cm(0)
                cell.paragraphs[0].add_run(text).bold = True

            for item in sorted(approvals, key=lambda value: value.get("order", 0)):
                row = table.add_row().cells
                row[0].text = str(item.get("order", ""))
                row[1].text = str(item.get("department", ""))
                row[2].text = str(item.get("term_days", ""))
                row[3].text = ""

        if related_clauses:
            document.add_paragraph()
            note = document.add_paragraph()
            note.paragraph_format.first_line_indent = Cm(0)
            note_run = note.add_run(
                "Документ подготовлен в соответствии с требованиями "
                f"{', '.join(related_clauses)} Инструкции по документообороту. "
                "Проект требует проверки перед направлением на согласование."
            )
            note_run.italic = True
            note_run.font.size = Pt(10)

        document.save(str(output_path))

    # ─────────────────── Публичный интерфейс ───────────────────
    async def generate(
        self,
        *,
        template_code: str,
        template_name: str,
        fields_schema: list[dict[str, Any]],
        values: dict[str, Any],
        template_file: Path | None = None,
        approvals: list[dict[str, Any]] | None = None,
        related_clauses: list[str] | None = None,
        author_name: str = "",
        author_position: str = "",
        department: str = "",
    ) -> tuple[Path, int]:
        """Формирует черновик документа.

        Args:
            template_code: Код шаблона.
            template_name: Название шаблона.
            fields_schema: Схема полей.
            values: Собранные значения полей.
            template_file: Путь к файлу бланка, если он есть.
            approvals: Маршрут согласования.
            related_clauses: Пункты инструкции с требованиями.
            author_name: ФИО автора для подписи.
            author_position: Должность автора.
            department: Подразделение автора.

        Returns:
            Кортеж ``(путь к файлу, размер в байтах)``.

        Raises:
            DocumentParseError: Не удалось сформировать документ.
        """
        context: dict[str, Any] = {
            **values,
            "doc_date": format_date_ru(),
            "author_name": author_name or values.get("author_name", ""),
            "author_position": author_position or values.get("author_position", ""),
            "department": department or values.get("department", ""),
        }

        filename = (
            f"{safe_filename(template_name)}_"
            f"{datetime.now(UTC).strftime('%Y%m%d_%H%M%S')}_"
            f"{uuid.uuid4().hex[:6]}.docx"
        )
        output_path = self.output_dir / filename

        if template_file is not None and template_file.exists():
            await asyncio.to_thread(
                self._render_from_template_sync, template_file, context, output_path
            )
            mode = "template"
        else:
            await asyncio.to_thread(
                self._build_structured_sync,
                title=template_name,
                fields_schema=fields_schema,
                values=context,
                approvals=approvals or [],
                related_clauses=related_clauses or [],
                output_path=output_path,
            )
            mode = "structured"

        size = output_path.stat().st_size
        documents_generated_total.labels(template_code=template_code).inc()

        logger.info(
            "document_generated",
            template_code=template_code,
            mode=mode,
            size_bytes=size,
        )
        return output_path, size


_generator: DocxGenerator | None = None


def get_docx_generator() -> DocxGenerator:
    """Возвращает singleton-генератор документов.

    Returns:
        Общий экземпляр :class:`DocxGenerator`.
    """
    global _generator
    if _generator is None:
        _generator = DocxGenerator()
    return _generator
